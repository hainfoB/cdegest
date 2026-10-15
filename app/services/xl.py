"""Classeurs Excel : en-têtes à la charte, sens de lecture selon la langue (arabe = droite à gauche)."""
import io

from flask import send_file

from ..i18n import get_lang

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def book():
    from openpyxl import Workbook
    return Workbook()


def sheet(wb, title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = get_lang() == "ar"
    return ws


def header(ws, row=1):
    from openpyxl.styles import Font, PatternFill
    for c in ws[row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D1B2A")


def widths(ws, *w):
    from openpyxl.utils import get_column_letter
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def send(wb, name):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=name, mimetype=XLSX)
