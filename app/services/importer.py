"""Lecture souple du fichier Excel des inscrits extrait de fp.cde.dz (en-têtes en arabe, français ou anglais)."""
import re
import unicodedata
from datetime import date, datetime

from ..models import READINESS_Q

# champ -> mots-clés reconnus dans l'en-tête (comparaison sans accents ni casse)
FIELDS = {
    "last_name": ["nom de famille", "nom", "اللقب", "last name", "family name", "surname"],
    "first_name": ["prenom", "الاسم", "first name", "given name"],
    "phone": ["telephone", "tel", "mobile", "portable", "الهاتف", "رقم الهاتف", "phone"],
    "email": ["e-mail", "email", "courriel", "البريد", "mail"],
    "birth_date": ["date de naissance", "naissance", "تاريخ الميلاد", "تاريخ الازدياد", "birth"],
    "sex": ["sexe", "genre", "الجنس", "sex", "gender"],
    "wilaya": ["wilaya", "الولاية", "ولاية"],
    "specialty": ["specialite", "التخصص", "تخصص", "specialty", "filiere", "الشعبة"],
    "diploma_year": ["annee du diplome", "annee", "سنة التخرج", "سنة", "year"],
    "origin": ["etablissement", "institut", "مؤسسة", "المؤسسة", "institution", "school"],
    "idea": ["idee", "projet", "فكرة", "المشروع", "idea", "project"],
    "pstatus": ["statut", "situation", "الوضعية", "الصفة", "status"],
}
# questions de l'étape 1 de fp.cde.dz : un fragment distinctif suffit
READY_KEYS = {
    "discover": ["التعرف على عالم المقاولاتية", "decouvrir", "discover"],
    "skills": ["تعلم مهارات", "apprendre", "learn"],
    "create_soon": ["مؤسسة عن قريب", "prochainement", "soon"],
    "project_ready": ["تجسيد المشروع", "concretiser", "already studied"],
    "business": ["مجال الأعمال", "domaine des affaires", "business"],
    "initiative": ["روح المبادرة", "esprit d'initiative", "initiative"],
}
YES = {"1", "oui", "yes", "نعم", "o", "y", "x", "vrai", "true"}


def norm(v):
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode() if re.search(r"[A-Za-z]", str(v or "")) else str(v or "")
    return re.sub(r"\s+", " ", s.replace("_", " ")).strip().lower()


def map_headers(headers):
    """Index de colonne pour chaque champ reconnu ; les questions d'abord (elles contiennent des mots comme « projet »)."""
    cols, used = {}, set()
    hs = [norm(h) for h in headers]
    for k, keys in READY_KEYS.items():
        for i, h in enumerate(hs):
            if i not in used and any(norm(x) in h for x in keys) and len(h) > 12:
                cols["q_" + k] = i
                used.add(i)
                break
    for f, keys in FIELDS.items():
        for kw in keys:                      # le mot-clé le plus précis d'abord
            pat = re.compile(r"(?<!\w)" + re.escape(norm(kw)) + r"(?!\w)")
            hit = next((i for i, h in enumerate(hs) if i not in used and pat.search(h)), None)
            if hit is not None:
                cols[f] = hit
                used.add(hit)
                break
    return cols


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def _sex(v):
    s = norm(v)
    if s in ("f", "femme", "feminin", "انثى", "أنثى", "female"):
        return "F"
    if s in ("m", "h", "homme", "masculin", "ذكر", "male"):
        return "M"
    return None


def read_rows(file_storage):
    from openpyxl import load_workbook
    wb = load_workbook(file_storage, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    headers = []
    for row in it:                            # première ligne non vide = en-têtes
        if row and sum(1 for c in row if c not in (None, "")) >= 2:
            headers = list(row)
            break
    cols = map_headers(headers)
    out = []
    for n, row in enumerate(it, 2):
        if not row or all(c in (None, "") for c in row):
            continue
        g = lambda f: row[cols[f]] if f in cols and cols[f] < len(row) else None  # noqa: E731
        phone = "".join(ch for ch in str(g("phone") or "") if ch.isdigit() or ch == "+")[:20]
        if phone and not phone.startswith(("0", "+")) and len(phone) == 9:
            phone = "0" + phone                # Excel enlève souvent le 0 initial
        y = re.search(r"(19|20)\d{2}", str(g("diploma_year") or ""))
        r = {
            "line": n,
            "last_name": str(g("last_name") or "").strip()[:80],
            "first_name": str(g("first_name") or "").strip()[:80],
            "phone": phone,
            "email": str(g("email") or "").strip()[:150] or None,
            "birth_date": (_date(g("birth_date")) or "") and _date(g("birth_date")).isoformat(),
            "sex": _sex(g("sex")),
            "wilaya": str(g("wilaya") or "").strip(),
            "specialty": str(g("specialty") or "").strip()[:120],
            "diploma_year": int(y.group(0)) if y else None,
            "origin": str(g("origin") or "").strip()[:150],
            "idea": str(g("idea") or "").strip()[:1000] or None,
            "pstatus": "stagiaire" if "stag" in norm(g("pstatus")) or "متربص" in str(g("pstatus") or "") else "diplome",
            "answers": ",".join(f"{q}={1 if norm(g('q_' + q)) in YES else 0}" for q in READINESS_Q),
        }
        r["errors"] = [e for e, bad in (("name", not (r["last_name"] and r["first_name"])),
                                        ("phone", len(r["phone"].lstrip("+")) < 9)) if bad]
        out.append(r)
    return cols, out


def template_headers(lang):
    fr = ["Nom", "Prénom", "Téléphone", "E-mail", "Date de naissance", "Sexe", "Wilaya", "Spécialité",
          "Année du diplôme", "Établissement", "Idée de projet"]
    ar = ["اللقب", "الاسم", "رقم الهاتف", "البريد الإلكتروني", "تاريخ الميلاد", "الجنس", "الولاية", "التخصص",
          "سنة التخرج", "المؤسسة", "فكرة المشروع"]
    q_ar = ["أرغب في التكوين فقط من أجل التعرف على عالم المقاولاتية", "أرغب في التكوين من أجل تعلم مهارات إنشاء مؤسسة استغلها مستقبلاً",
            "أرغب في التكوين من أجل إنشاء مؤسسة عن قريب", "أرغب في التكوين من أجل تجسيد المشروع الذي أنهيت دراسته",
            "أحب مجال الأعمال وأرغب في الولوج إليه", "لدي روح المبادرة"]
    return (ar if lang == "ar" else fr) + q_ar
