"""Rapports semestriels : institut → wilaya (consolidation, validation DFEP) → MFEP. Chaque rapport transmis est figé."""
import io
import json

from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user

from ..extensions import db
from ..i18n import t
from ..models import AuditLog, Institute, Report, Session, Wilaya
from ..security import ADMIN, COORD, DFEP, MFEP, RESP, roles_required, visible_wilayas
from ..services import kpi as K
from ..services.period import Period, recent_years
from ..utils import utcnow

bp = Blueprint("reports", __name__, url_prefix="/rapports")


def _snapshot_institute(inst, period):
    values = K.kpis_for_institutes([inst.id], period)
    tg = K.targets(period, inst.wilaya_id, inst.id)
    per, status = K.evaluate(values, tg, period)
    sessions = []
    for s in Session.query.filter(Session.institute_id == inst.id,
                                  Session.starts_on.between(period.start, period.end),
                                  Session.status != "draft").order_by(Session.starts_on):
        conf = [e for e in s.enrollments if e.status == "confirmed"]
        sessions.append({"title": s.title, "status": s.status, "groups": len(s.groups), "confirmed": len(conf),
                         "dropouts": sum(1 for e in conf if e.dropout),
                         "certificates": sum(1 for e in conf if e.cert_no), "nesda": s.nesda_status,
                         "teachers": len(s.teacher_ids())})
    return {"values": values, "targets": tg, "per": per, "status": status, "sessions": sessions,
            "institute": {"code": inst.code, "fr": inst.name_fr, "ar": inst.name_ar},
            "generated_at": utcnow().isoformat(timespec="minutes")}


def _snapshot_wilaya(w, period):
    insts = [i for i in w.institutes if i.active]
    rows = K.institute_table(insts, period)
    values = K.finalize(K.merge(K.raw_by_institute([i.id for i in insts], period.start, period.end).values()))
    tg = K.targets(period, w.id)
    per, status = K.evaluate(values, tg, period)
    sent = {r.institute_id: r.status for r in Report.query.filter_by(level="institute", wilaya_id=w.id,
                                                                      year=period.year, semester=period.semester)}
    return {"values": values, "targets": tg, "per": per, "status": status,
            "institutes": [{"code": r["institute"].code, "fr": r["institute"].name_fr, "ar": r["institute"].name_ar,
                            "values": r["values"], "status": r["status"],
                            "report": sent.get(r["institute"].id, "none")} for r in rows],
            "wilaya": {"code": w.code, "fr": w.name_fr, "ar": w.name_ar},
            "generated_at": utcnow().isoformat(timespec="minutes")}


def _period():
    p = Period.from_request(request.values)
    if p.semester == 0:          # reports are per semester
        p = Period(p.year, "S1")
    return p


@bp.route("/", methods=["GET", "POST"])
@roles_required(ADMIN, MFEP, DFEP, COORD, RESP)
def index():
    me = current_user
    period = _period()
    if request.method == "POST":
        act = request.form.get("action")
        if me.role == RESP or (me.role == ADMIN and request.form.get("institute")):
            inst = me.institute if me.role == RESP else db.session.get(Institute, int(request.form["institute"]))
            r = Report.query.filter_by(level="institute", institute_id=inst.id, year=period.year,
                                       semester=period.semester).first()
            if r and r.status == "transmitted":
                abort(403)
            r = r or Report(level="institute", institute_id=inst.id, wilaya_id=inst.wilaya_id, year=period.year,
                            semester=period.semester)
            r.snapshot = json.dumps(_snapshot_institute(inst, period))
            r.generated_at = utcnow()
            r.comment = (request.form.get("comment") or r.comment or "")[:2000]
            if act == "transmit":
                r.status, r.transmitted_by_id, r.transmitted_at = "transmitted", me.id, utcnow()
            db.session.add(r)
        elif me.role in (DFEP, COORD, ADMIN):
            w = me.wilaya if me.role != ADMIN else Wilaya.query.filter_by(code=request.form.get("w")).first_or_404()
            r = Report.query.filter_by(level="wilaya", institute_id=None, wilaya_id=w.id, year=period.year,
                                       semester=period.semester).first()
            if r and r.status == "transmitted":
                abort(403)
            if act == "transmit" and me.role not in (DFEP, ADMIN):
                abort(403)              # only the DFEP validates and transmits to the MFEP
            r = r or Report(level="wilaya", wilaya_id=w.id, year=period.year, semester=period.semester)
            r.snapshot = json.dumps(_snapshot_wilaya(w, period))
            r.generated_at = utcnow()
            r.comment = (request.form.get("comment") or r.comment or "")[:2000]
            if act == "transmit":
                r.status, r.transmitted_by_id, r.transmitted_at = "transmitted", me.id, utcnow()
            db.session.add(r)
        else:
            abort(403)
        AuditLog.log(f"report.{act}", details=f"{period.year} S{period.semester}")
        db.session.commit()
        flash(t("rep.transmitted") if act == "transmit" else t("rep.generated"), "ok")
        return redirect(url_for("reports.index", y=period.year, p=period.code))

    ctx = {"period": period, "years": recent_years(3)}
    if me.role == RESP:
        ctx["mine"] = Report.query.filter_by(level="institute", institute_id=me.institute_id, year=period.year,
                                             semester=period.semester).first()
        ctx["view"] = "institute"
    elif me.role in (DFEP, COORD):
        ctx["view"] = "wilaya"
        ctx["insts"] = [i for i in me.wilaya.institutes if i.active]
        ctx["inst_reports"] = {r.institute_id: r for r in Report.query.filter_by(
            level="institute", wilaya_id=me.wilaya_id, year=period.year, semester=period.semester)}
        ctx["mine"] = Report.query.filter_by(level="wilaya", wilaya_id=me.wilaya_id, year=period.year,
                                             semester=period.semester).first()
    else:
        ctx["view"] = "national"
        ctx["wilayas"] = visible_wilayas(me).all()
        ctx["w_reports"] = {r.wilaya_id: r for r in Report.query.filter_by(
            level="wilaya", year=period.year, semester=period.semester, status="transmitted")}
    return render_template("reports/index.html", **ctx)


def _get(rid):
    r = db.session.get(Report, rid) or abort(404)
    me = current_user
    if me.role in (ADMIN,):
        return r
    if me.role == MFEP and r.level == "wilaya" and r.status == "transmitted":
        return r
    if me.role in (DFEP, COORD) and r.wilaya_id == me.wilaya_id:
        return r
    if me.role == RESP and r.institute_id == me.institute_id:
        return r
    abort(403)


@bp.route("/<int:rid>")
@roles_required(ADMIN, MFEP, DFEP, COORD, RESP)
def view(rid):
    r = _get(rid)
    return render_template("reports/view.html", r=r, d=r.data(), families=K.KPI_FAMILIES, fmt=K.fmt)


@bp.route("/<int:rid>.xlsx")
@roles_required(ADMIN, MFEP, DFEP, COORD, RESP)
def export(rid):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    r = _get(rid)
    d = r.data()
    wb = Workbook()
    ws = wb.active
    ws.title = "Indicateurs"
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0D1B2A")
    scope = d.get("institute") or d.get("wilaya") or {}
    ws.append([f"CDEgest — rapport {'institut' if r.level == 'institute' else 'wilaya'} — {scope.get('fr', '')}"])
    ws.append([f"Année professionnelle {r.year} — semestre {r.semester} — {r.status}"])
    ws.append([])
    ws.append(["Indicateur", "Réalisé", "Cible", "Statut"])
    for c in ws[4]:
        c.font, c.fill = hdr, fill
    for k in K.KPIS:
        ws.append([k, d["values"].get(k), d["targets"].get(k), d["per"].get(k)])
    ws.column_dimensions["A"].width = 22
    if d.get("institutes"):
        ws2 = wb.create_sheet("Instituts")
        ws2.append(["Code", "Institut", "Statut objectifs", "Rapport"] + K.KPIS)
        for c in ws2[1]:
            c.font, c.fill = hdr, fill
        for i in d["institutes"]:
            ws2.append([i["code"], i["fr"], i["status"], i["report"]] + [i["values"].get(k) for k in K.KPIS])
    if d.get("sessions"):
        ws3 = wb.create_sheet("Sessions")
        ws3.append(["Session", "Statut", "Groupes", "Inscrits", "Abandons", "Attestations", "NESDA", "Enseignants"])
        for c in ws3[1]:
            c.font, c.fill = hdr, fill
        for s in d["sessions"]:
            ws3.append([s["title"], s["status"], s["groups"], s["confirmed"], s["dropouts"], s["certificates"],
                        s["nesda"], s["teachers"]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"rapport-{r.level}-{r.year}-S{r.semester}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
